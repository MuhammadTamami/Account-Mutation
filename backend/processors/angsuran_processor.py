"""
Angsuran Processor - Extract & Generate Proyeksi KOP
Processes batch Excel files containing loan installment schedules
"""

import pandas as pd
import numpy as np
from datetime import datetime
import os
import logging

logger = logging.getLogger(__name__)


def process_angsuran_batch(file_paths, filter_month=None, filter_year=None):
    """
    Process batch of angsuran Excel files
    
    Args:
        file_paths: List of file paths to process
        filter_month: Optional month filter (1-12)
        filter_year: Optional year filter (e.g. 2026)
    
    Returns:
        dict: {
            'summary': DataFrame with aggregated data,
            'details': List of DataFrames per file,
            'stats': Statistics
        }
    """
    all_details = []
    errors = []
    
    for file_path in file_paths:
        try:
            logger.info(f"Processing {os.path.basename(file_path)}")
            df = extract_angsuran_data(file_path)
            
            if df is not None and len(df) > 0:
                all_details.append({
                    'filename': os.path.basename(file_path),
                    'data': df
                })
        except Exception as e:
            logger.error(f"Error processing {file_path}: {e}")
            errors.append({
                'file': os.path.basename(file_path),
                'error': str(e)
            })
    
    if not all_details:
        raise Exception(f"No valid data extracted from files. Errors: {errors}")
    
    # Generate summary (aggregate by period)
    summary_df = generate_summary(all_details, filter_month, filter_year)
    
    # Calculate statistics
    stats = calculate_statistics(summary_df)
    
    return {
        'summary': summary_df,
        'details': all_details,
        'stats': stats,
        'errors': errors
    }


def extract_angsuran_data(file_path):
    """
    Extract loan installment data from Excel file
    
    Expected structure:
    - Row with "No Kontrak :" contains contract number
    - Row with "No.", "Tanggal", "Total Bayar", "Pokok", "Margin", "Baki Debet" headers
    - Data rows below headers
    """
    try:
        # Read Excel file (try different sheet names)
        xl = pd.ExcelFile(file_path)
        sheet_name = xl.sheet_names[0]  # Use first sheet
        
        df_raw = pd.read_excel(file_path, sheet_name=sheet_name, header=None)
        
        # Find contract number (row with "No Kontrak :")
        contract_row = df_raw[df_raw.iloc[:, 0].astype(str).str.contains('No Kontrak', case=False, na=False)]
        
        if len(contract_row) == 0:
            raise Exception("Cannot find 'No Kontrak' row")
        
        contract_number = str(contract_row.iloc[0, 1]).strip()
        
        # Find plafon (initial amount - first row after header with valid Baki Debet)
        plafon = None
        
        # Find header row (contains "No.", "Tanggal", "Total Bayar", etc.)
        header_row = None
        for idx, row in df_raw.iterrows():
            row_str = ' '.join(row.astype(str).str.lower().tolist())
            if 'tanggal' in row_str and 'total bayar' in row_str and 'pokok' in row_str:
                header_row = idx
                break
        
        if header_row is None:
            raise Exception("Cannot find header row with 'Tanggal', 'Total Bayar', 'Pokok'")
        
        # Read data starting from header
        df = pd.read_excel(file_path, sheet_name=sheet_name, header=header_row)
        
        # Clean column names
        df.columns = df.columns.astype(str).str.strip()
        
        # Find relevant columns (case-insensitive matching)
        col_map = {}
        for col in df.columns:
            col_lower = col.lower()
            if 'tanggal' in col_lower and 'no' not in col_lower:
                col_map['Tanggal'] = col
            elif 'total' in col_lower and 'bayar' in col_lower:
                col_map['Total Bayar'] = col
            elif col_lower == 'pokok' or 'pokok' in col_lower:
                col_map['Pokok'] = col
            elif 'margin' in col_lower:
                col_map['Margin'] = col
            elif 'baki' in col_lower and 'debet' in col_lower:
                col_map['Baki Debet'] = col
        
        if 'Tanggal' not in col_map:
            raise Exception("Cannot find 'Tanggal' column")
        
        # Select and rename columns
        required_cols = ['Tanggal', 'Total Bayar', 'Pokok', 'Margin', 'Baki Debet']
        df_clean = df[[col_map.get(c, c) for c in required_cols if c in col_map]].copy()
        df_clean.columns = [c for c in required_cols if c in col_map]
        
        # Convert Tanggal to datetime
        df_clean['Tanggal'] = pd.to_datetime(df_clean['Tanggal'], errors='coerce')
        
        # Fix year if it's clearly wrong (before 2000)
        # Excel sometimes stores dates with wrong years
        mask_wrong_year = df_clean['Tanggal'].dt.year < 2000
        if mask_wrong_year.any():
            # Add 96 years to fix 1930 -> 2026
            df_clean.loc[mask_wrong_year, 'Tanggal'] = df_clean.loc[mask_wrong_year, 'Tanggal'] + pd.DateOffset(years=96)
        
        # Remove rows with invalid dates
        df_clean = df_clean[df_clean['Tanggal'].notna()].copy()
        
        if len(df_clean) == 0:
            raise Exception("No valid data rows found")
        
        # Convert numeric columns
        for col in ['Total Bayar', 'Pokok', 'Margin', 'Baki Debet']:
            if col in df_clean.columns:
                df_clean[col] = pd.to_numeric(df_clean[col], errors='coerce').fillna(0)
        
        # Add contract info
        df_clean['No Kontrak'] = contract_number
        
        # Plafon is the FIRST Baki Debet (initial loan amount)
        if 'Baki Debet' in df_clean.columns and len(df_clean) > 0:
            plafon = df_clean['Baki Debet'].iloc[0]
        else:
            plafon = 0
        
        df_clean['Plafon'] = plafon
        
        # Outstanding is same as Baki Debet for each row
        df_clean['Outstanding'] = df_clean['Baki Debet'] if 'Baki Debet' in df_clean.columns else 0
        
        return df_clean
    
    except Exception as e:
        logger.error(f"Error extracting data from {file_path}: {e}")
        raise


def generate_summary(details_list, filter_month=None, filter_year=None):
    """
    Generate summary DataFrame aggregated by period
    
    Args:
        details_list: List of dicts with 'filename' and 'data' keys
        filter_month: Optional month filter (1-12)
        filter_year: Optional year filter (e.g. 2026)
    
    Returns:
        DataFrame with columns: No, Fasilitas Pembiayaan, Plafon, Outstanding, Periode, 
                               Porsi Pokok, Porsi Margin, Total
    """
    summary_data = []
    
    for idx, detail in enumerate(details_list, 1):
        df = detail['data']
        contract = df['No Kontrak'].iloc[0]
        plafon = df['Plafon'].iloc[0]
        
        # Apply filter
        if filter_month or filter_year:
            df_filtered = df.copy()
            
            if filter_year:
                df_filtered = df_filtered[df_filtered['Tanggal'].dt.year == int(filter_year)]
            
            if filter_month:
                df_filtered = df_filtered[df_filtered['Tanggal'].dt.month == int(filter_month)]
            
            if len(df_filtered) == 0:
                logger.warning(f"No data for contract {contract} in period {filter_month}/{filter_year}")
                continue
        else:
            # Use all data
            df_filtered = df
        
        # Get first date in filtered period - ALWAYS use day 1 of the month
        if filter_month and filter_year:
            # Set periode to first day of filtered month
            periode = pd.Timestamp(year=int(filter_year), month=int(filter_month), day=1)
        else:
            # No filter: use first transaction date but set to day 1
            first_date = df_filtered['Tanggal'].min()
            periode = pd.Timestamp(year=first_date.year, month=first_date.month, day=1)
        
        # Get FIRST transaction installment in filtered period (not SUM!)
        # Ambil SATU transaksi pertama dari bulan filter, bukan dijumlahkan
        first_transaction = df_filtered.iloc[0]
        
        porsi_pokok = first_transaction.get('Pokok', 0) if 'Pokok' in df_filtered.columns else 0
        porsi_margin = first_transaction.get('Margin', 0) if 'Margin' in df_filtered.columns else 0
        total = first_transaction.get('Total Bayar', 0) if 'Total Bayar' in df_filtered.columns else 0
        
        # Outstanding: Ambil Baki Debet dari bulan SEBELUM filter
        # Jika filter Oktober, maka Outstanding ambil dari Baki Debet PERTAMA di September
        # Porsi Pokok & Margin tetap ambil dari Oktober (sesuai filter - transaksi pertama)
        
        if filter_month and filter_year:
            # Calculate previous month
            target_year = int(filter_year)
            target_month = int(filter_month)
            
            if target_month == 1:
                prev_month = 12
                prev_year = target_year - 1
            else:
                prev_month = target_month - 1
                prev_year = target_year
            
            # Get transactions from previous month
            transactions_prev_month = df[
                (df['Tanggal'].dt.year == prev_year) & 
                (df['Tanggal'].dt.month == prev_month)
            ]
            
            if len(transactions_prev_month) > 0:
                # Use the FIRST Baki Debet from previous month (bukan last!)
                if 'Baki Debet' in transactions_prev_month.columns:
                    outstanding = transactions_prev_month['Baki Debet'].iloc[0]
                else:
                    outstanding = plafon
            else:
                # No transactions in previous month, look for any transaction before filtered period
                first_period_date = df_filtered['Tanggal'].min()
                transactions_before = df[df['Tanggal'] < first_period_date]
                
                if len(transactions_before) > 0 and 'Baki Debet' in transactions_before.columns:
                    outstanding = transactions_before['Baki Debet'].iloc[0]
                else:
                    outstanding = plafon
        else:
            # No filter: use Baki Debet from first transaction in dataset
            first_period_date = df_filtered['Tanggal'].min()
            transactions_before = df[df['Tanggal'] < first_period_date]
            
            if len(transactions_before) > 0 and 'Baki Debet' in transactions_before.columns:
                outstanding = transactions_before['Baki Debet'].iloc[0]
            else:
                outstanding = plafon
        
        summary_data.append({
            'No': idx,
            'Fasilitas Pembiayaan': contract,
            'Plafon': plafon,
            'Outstanding': outstanding,
            'Periode': periode,
            'Porsi Pokok': porsi_pokok,
            'Porsi Margin': porsi_margin,
            'Total': total
        })
    
    if not summary_data:
        raise Exception("No data matches the filter criteria")
    
    summary_df = pd.DataFrame(summary_data)
    
    return summary_df


def calculate_statistics(summary_df):
    """Calculate statistics from summary"""
    stats = {
        'total_contracts': len(summary_df),
        'total_plafon': summary_df['Plafon'].sum(),
        'total_outstanding': summary_df['Outstanding'].sum(),
        'total_porsi_pokok': summary_df['Porsi Pokok'].sum(),
        'total_porsi_margin': summary_df['Porsi Margin'].sum(),
        'total_angsuran': summary_df['Total'].sum()
    }
    
    return stats


def export_to_excel_angsuran(summary_df, stats, output_path, sheet_name='Proyeksi'):
    """
    Export summary to Excel with professional formatting
    
    Features:
    - Custom currency format: _(Rp* #,##0.00_);_(Rp* (#,##0.00);_(Rp* "-"??_);_(@_)
    - Fasilitas Pembiayaan: Left align
    - Numeric columns (Plafon, Outstanding, Porsi Pokok, Porsi Margin, Total): Right align
    - Total column: Uses formula (=F+G) not static value
    """
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
    from openpyxl.styles.numbers import FORMAT_NUMBER_00
    
    wb = Workbook()
    ws = wb.active
    ws.title = sheet_name
    
    # Define styles
    header_fill = PatternFill(start_color="D9D9D9", end_color="D9D9D9", fill_type="solid")
    header_font = Font(bold=True, size=11, name='Calibri')
    header_alignment = Alignment(horizontal='center', vertical='center')
    
    thin_border = Border(
        left=Side(style='thin', color='000000'),
        right=Side(style='thin', color='000000'),
        top=Side(style='thin', color='000000'),
        bottom=Side(style='thin', color='000000')
    )
    
    # Cell alignment
    left_align = Alignment(horizontal='left', vertical='center')
    center_align = Alignment(horizontal='center', vertical='center')
    right_align = Alignment(horizontal='right', vertical='center')
    
    # Data font
    data_font = Font(size=11, name='Calibri')
    
    # Custom currency format (Rp without text prefix, using Excel number format)
    currency_format = '_(Rp* #,##0.00_);_(Rp* (#,##0.00);_(Rp* "-"??_);_(@_)'
    
    # Write headers (row 3)
    headers = ['No', 'Fasilitas Pembiayaan', 'Plafon', 'Outstanding', 'Periode', 
               'Porsi Pokok', 'Porsi Margin', 'Total']
    
    for col_idx, header in enumerate(headers, 1):
        cell = ws.cell(row=3, column=col_idx, value=header)
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = header_alignment
        cell.border = thin_border
    
    # Write data (starting row 4)
    start_row = 4
    for idx, row_data in summary_df.iterrows():
        row_num = start_row + idx
        
        # Column A: No
        cell = ws.cell(row=row_num, column=1, value=int(row_data['No']))
        cell.alignment = center_align
        cell.font = data_font
        cell.border = thin_border
        
        # Column B: Fasilitas Pembiayaan (LEFT ALIGN)
        cell = ws.cell(row=row_num, column=2, value=str(row_data['Fasilitas Pembiayaan']))
        cell.alignment = left_align  # LEFT ALIGN
        cell.font = data_font
        cell.border = thin_border
        
        # Column C: Plafon (RIGHT ALIGN with custom currency format)
        cell = ws.cell(row=row_num, column=3, value=float(row_data['Plafon']))
        cell.number_format = currency_format
        cell.alignment = right_align  # RIGHT ALIGN
        cell.font = data_font
        cell.border = thin_border
        
        # Column D: Outstanding (RIGHT ALIGN with custom currency format)
        cell = ws.cell(row=row_num, column=4, value=float(row_data['Outstanding']))
        cell.number_format = currency_format
        cell.alignment = right_align  # RIGHT ALIGN
        cell.font = data_font
        cell.border = thin_border
        
        # Column E: Periode (format: Oct-26)
        periode_val = row_data['Periode']
        if pd.notna(periode_val):
            periode_dt = pd.to_datetime(periode_val)
            periode_str = periode_dt.strftime('%b-%y')
        else:
            periode_str = ''
        
        cell = ws.cell(row=row_num, column=5, value=periode_str)
        cell.alignment = center_align
        cell.font = data_font
        cell.border = thin_border
        
        # Column F: Porsi Pokok (RIGHT ALIGN with custom currency format)
        cell = ws.cell(row=row_num, column=6, value=float(row_data['Porsi Pokok']))
        cell.number_format = currency_format
        cell.alignment = right_align  # RIGHT ALIGN
        cell.font = data_font
        cell.border = thin_border
        
        # Column G: Porsi Margin (RIGHT ALIGN with custom currency format)
        porsi_margin = row_data['Porsi Margin']
        cell = ws.cell(row=row_num, column=7)
        if porsi_margin == 0 or pd.isna(porsi_margin):
            cell.value = 0
        else:
            cell.value = float(porsi_margin)
        cell.number_format = currency_format
        cell.alignment = right_align  # RIGHT ALIGN
        cell.font = data_font
        cell.border = thin_border
        
        # Column H: Total (FORMULA: =F+G, not static value)
        cell = ws.cell(row=row_num, column=8)
        cell.value = f'=F{row_num}+G{row_num}'  # FORMULA instead of static value
        cell.number_format = currency_format
        cell.alignment = right_align  # RIGHT ALIGN
        cell.font = data_font
        cell.border = thin_border
    
    # Add Total Row
    total_row = start_row + len(summary_df) + 1  # +1 for blank row
    
    # Merge cells for "Total" label (columns A-G)
    ws.merge_cells(f'A{total_row}:G{total_row}')
    cell_total_label = ws.cell(row=total_row, column=1, value='Total')
    cell_total_label.alignment = center_align
    cell_total_label.font = Font(bold=True, size=12, name='Calibri')
    cell_total_label.border = thin_border
    
    # Total value in column H (SUM formula)
    cell_total_value = ws.cell(row=total_row, column=8)
    cell_total_value.value = f'=SUM(H{start_row}:H{start_row + len(summary_df) - 1})'
    cell_total_value.number_format = currency_format
    cell_total_value.alignment = right_align
    cell_total_value.font = Font(bold=True, size=12, name='Calibri')
    cell_total_value.border = thin_border
    
    # Set column widths
    ws.column_dimensions['A'].width = 5    # No
    ws.column_dimensions['B'].width = 25   # Fasilitas Pembiayaan (wider for left-aligned text)
    ws.column_dimensions['C'].width = 22   # Plafon
    ws.column_dimensions['D'].width = 22   # Outstanding
    ws.column_dimensions['E'].width = 12   # Periode
    ws.column_dimensions['F'].width = 20   # Porsi Pokok
    ws.column_dimensions['G'].width = 20   # Porsi Margin
    ws.column_dimensions['H'].width = 22   # Total
    
    # Save workbook
    wb.save(output_path)
    logger.info(f"Exported Angsuran Excel: {output_path}")
