# -*- coding: utf-8 -*-
"""
Verify Excel format
"""
from openpyxl import load_workbook

# Load the generated Excel file
wb = load_workbook('test_ideb_with_total.xlsx')
ws = wb.active

print("Excel Content Preview:")
print("=" * 100)

# Show cell B4 (Debitur name)
print(f"B4 (Debitur Name): {ws['B4'].value}")
print()

# Show header row (row 5, starting from B)
print("Row 5 (Headers):")
headers = []
for col in range(2, 12):  # B to K (columns 2-11)
    cell = ws.cell(row=5, column=col)
    headers.append(cell.value)
    print(f"  Column {chr(64+col)}: {cell.value}")
print()

# Show data rows (starting from row 6)
print("Data Rows:")
for row_num in range(6, 10):  # Rows 6-9 (3 data + 1 total)
    row_data = []
    for col in range(2, 12):  # B to K
        cell = ws.cell(row=row_num, column=col)
        value = cell.value
        # Check if it's a date
        if hasattr(value, 'strftime'):
            value = value.strftime('%d-%b-%y')
        row_data.append(str(value) if value is not None else '')
    print(f"  Row {row_num}: {' | '.join(row_data)}")

print("=" * 100)

# Check formatting
print("\nFormatting Check:")
print(f"B4 Font Bold: {ws['B4'].font.bold}")
print(f"B5 Fill Color: {ws['B5'].fill.start_color.rgb if ws['B5'].fill.start_color else 'None'}")
print(f"Total Row (B9) Fill Color: {ws['B9'].fill.start_color.rgb if ws['B9'].fill.start_color else 'None'}")
print(f"E6 (Yield) Number Format: {ws['E6'].number_format}")
print(f"G6 (Tanggal Pencairan) Number Format: {ws['G6'].number_format}")

wb.close()
