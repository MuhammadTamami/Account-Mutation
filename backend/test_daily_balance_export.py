"""
Test daily balance Excel export
"""
import pandas as pd
from openpyxl.styles import Font, PatternFill, Border, Side, Alignment
import os

# Simulate daily balance data
test_data = {
    'Tanggal': [1, 2, 3, 8, 9, 10],
    'Saldo': ['2.108.916.635,31', '2.031.524.682,31', '2.053.637.494,31', 
              '2.133.034.929,31', '2.119.692.325,31', '2.138.619.325,31']
}

df = pd.DataFrame(test_data)

print("Test Data:")
print(df)
print()

# Create Excel file
output_path = os.path.join('outputs', 'test_daily_balance.xlsx')

with pd.ExcelWriter(output_path, engine='openpyxl') as writer:
    sheet_name = 'Saldo Harian'
    df.to_excel(writer, index=False, sheet_name=sheet_name)
    
    worksheet = writer.sheets[sheet_name]
    
    # Parse balance values
    def parse_amount(val):
        if pd.isna(val) or val == '' or val == '-':
            return 0.0
        val_str = str(val).strip().replace('Rp', '').replace(' ', '').strip()
        
        # Handle different formats
        if ',' in val_str and '.' in val_str:
            last_comma_pos = val_str.rfind(',')
            last_dot_pos = val_str.rfind('.')
            
            if last_dot_pos > last_comma_pos:
                # International: 1,234,567.89
                val_str = val_str.replace(',', '')
            else:
                # Indonesian: 1.234.567,89
                val_str = val_str.replace('.', '').replace(',', '.')
        elif ',' in val_str:
            val_str = val_str.replace('.', '').replace(',', '.')
        elif '.' in val_str:
            last_dot_pos = val_str.rfind('.')
            decimal_part_length = len(val_str) - last_dot_pos - 1
            if decimal_part_length == 2:
                val_str = val_str[:last_dot_pos].replace('.', '') + '.' + val_str[last_dot_pos+1:]
            else:
                val_str = val_str.replace('.', '')
        
        try:
            return float(val_str)
        except:
            print(f"Failed to parse: {val} -> {val_str}")
            return 0.0
    
    # Calculate statistics
    df_numeric = df.copy()
    df_numeric['Balance_numeric'] = df_numeric['Saldo'].apply(parse_amount)
    
    print("Parsed balances:")
    print(df_numeric[['Saldo', 'Balance_numeric']])
    print()
    
    total_balance = df_numeric['Balance_numeric'].sum()
    avg_balance = df_numeric['Balance_numeric'].mean()
    highest_balance = df_numeric['Balance_numeric'].max()
    lowest_balance = df_numeric['Balance_numeric'].min()
    
    print(f"Total: {total_balance:,.2f}")
    print(f"Average: {avg_balance:,.2f}")
    print(f"Highest: {highest_balance:,.2f}")
    print(f"Lowest: {lowest_balance:,.2f}")
    print()
    
    # Add statistics
    last_row = len(df) + 2  # +1 for header, +1 for blank row
    stats_start_row = last_row + 1
    
    # Statistics section
    worksheet.cell(row=stats_start_row, column=1, value='Total')
    worksheet.cell(row=stats_start_row, column=2, value=total_balance)
    worksheet.cell(row=stats_start_row, column=2).number_format = '#,##0.00'
    
    worksheet.cell(row=stats_start_row + 1, column=1, value='Rata-rata Pengendapan')
    worksheet.cell(row=stats_start_row + 1, column=2, value=avg_balance)
    worksheet.cell(row=stats_start_row + 1, column=2).number_format = '#,##0.00'
    
    worksheet.cell(row=stats_start_row + 2, column=1, value='Saldo Rata-rata')
    worksheet.cell(row=stats_start_row + 2, column=2, value=avg_balance)
    worksheet.cell(row=stats_start_row + 2, column=2).number_format = '#,##0.00'
    
    worksheet.cell(row=stats_start_row + 3, column=1, value='Saldo Tertinggi')
    worksheet.cell(row=stats_start_row + 3, column=2, value=highest_balance)
    worksheet.cell(row=stats_start_row + 3, column=2).number_format = '#,##0.00'
    
    worksheet.cell(row=stats_start_row + 4, column=1, value='Saldo Terendah')
    worksheet.cell(row=stats_start_row + 4, column=2, value=lowest_balance)
    worksheet.cell(row=stats_start_row + 4, column=2).number_format = '#,##0.00'
    
    # Mutation statistics table
    mutation_start_row = stats_start_row + 6
    
    # Header row with merged cells
    worksheet.merge_cells(start_row=mutation_start_row, start_column=1, end_row=mutation_start_row, end_column=3)
    header_cell = worksheet.cell(row=mutation_start_row, column=1)
    header_cell.value = 'Mutasi Debet'
    header_cell.fill = PatternFill(start_color="FFFFFF", end_color="FFFFFF", fill_type="solid")
    header_cell.font = Font(bold=True)
    header_cell.alignment = Alignment(horizontal='center')
    
    worksheet.merge_cells(start_row=mutation_start_row, start_column=4, end_row=mutation_start_row, end_column=6)
    header_cell2 = worksheet.cell(row=mutation_start_row, column=4)
    header_cell2.value = 'Mutasi kredit'
    header_cell2.fill = PatternFill(start_color="FFFFFF", end_color="FFFFFF", fill_type="solid")
    header_cell2.font = Font(bold=True)
    header_cell2.alignment = Alignment(horizontal='center')
    
    # Data rows
    mutation_rows = [
        ('Total Mutasi', '-', '-'),
        ('Adjusted', '-', '-'),
        ('Total Frekuensi', '-', '-'),
        ('Adjusted', '-', '-')
    ]
    
    thin_border = Border(
        left=Side(style='thin'),
        right=Side(style='thin'),
        top=Side(style='thin'),
        bottom=Side(style='thin')
    )
    
    for idx, (label, debit_val, credit_val) in enumerate(mutation_rows, start=1):
        row_num = mutation_start_row + idx
        worksheet.cell(row=row_num, column=1, value=label)
        worksheet.merge_cells(start_row=row_num, start_column=2, end_row=row_num, end_column=3)
        worksheet.cell(row=row_num, column=2, value=debit_val)
        worksheet.merge_cells(start_row=row_num, start_column=5, end_row=row_num, end_column=6)
        worksheet.cell(row=row_num, column=5, value=credit_val)
        
        # Add borders
        for col in range(1, 7):
            worksheet.cell(row=row_num, column=col).border = thin_border
    
    # Add borders to header
    for col in range(1, 7):
        worksheet.cell(row=mutation_start_row, column=col).border = thin_border
    
    # Column widths
    worksheet.column_dimensions['A'].width = 20
    worksheet.column_dimensions['B'].width = 20

print(f"\n✅ Test file created: {output_path}")
print("Please check the file to verify formatting.")
