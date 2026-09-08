"""
Test full download flow with Bank Kalsel data
"""
import sys
import os
sys.path.insert(0, os.path.join(os.path.dirname(__file__), 'processors'))

from processors.bank_kalsel_processor import process_bank_kalsel_file
import pandas as pd
from openpyxl.styles import Font, PatternFill, Border, Side, Alignment

# Process Bank Kalsel test file
pdf_path = r"..\test_data\bank_kalsel\1. Januari.pdf"

print("Processing Bank Kalsel PDF...")
df = process_bank_kalsel_file(pdf_path, 'pdf')

if df is None or df.empty:
    print("❌ Failed to process")
    sys.exit(1)

print(f"✅ Processed {len(df)} transactions")
print("\nFirst few transactions:")
print(df[['Date', 'Type', 'Amount', 'Balance']].head())

# Simulate daily balance extraction (like in app.py)
print("\n" + "="*60)
print("EXTRACTING DAILY BALANCE...")
print("="*60)

df['Date_parsed'] = pd.to_datetime(df['Date'], format='%d/%m/%Y', errors='coerce')
df['DateOnly'] = df['Date_parsed'].dt.date

# Group by date and get last transaction
daily_df = df.groupby('DateOnly').agg({
    'Date': 'first',
    'Date_parsed': 'first',
    'Balance': 'last'
}).reset_index()

# Extract day number
daily_df['Tanggal'] = daily_df['Date_parsed'].dt.day

# Prepare export dataframe
df_export = daily_df[['Tanggal', 'Balance']].copy()
df_export.columns = ['Tanggal', 'Saldo']

print(f"\n✅ Extracted {len(df_export)} daily balances")
print("\nDaily Balance Data:")
print(df_export)

# Create Excel file
output_path = os.path.join('outputs', 'test_bank_kalsel_daily.xlsx')

print("\n" + "="*60)
print("CREATING EXCEL FILE...")
print("="*60)

with pd.ExcelWriter(output_path, engine='openpyxl') as writer:
    sheet_name = 'Saldo Harian'
    df_export.to_excel(writer, index=False, sheet_name=sheet_name)
    
    worksheet = writer.sheets[sheet_name]
    
    # Parse balance values
    def parse_amount(val):
        if pd.isna(val) or val == '' or val == '-':
            return 0.0
        val_str = str(val).strip().replace('Rp', '').replace(' ', '').strip()
        
        # Handle different formats
        if ',' in val_str and '.' in val_str:
            last_comma_pos = val_str.rfind(',')
            last_dot_pos = val_str.rfind('.')
            
            if last_dot_pos > last_comma_pos:
                val_str = val_str.replace(',', '')
            else:
                val_str = val_str.replace('.', '').replace(',', '.')
        elif ',' in val_str:
            val_str = val_str.replace('.', '').replace(',', '.')
        elif '.' in val_str:
            last_dot_pos = val_str.rfind('.')
            decimal_part_length = len(val_str) - last_dot_pos - 1
            if decimal_part_length == 2:
                val_str = val_str[:last_dot_pos].replace('.', '') + '.' + val_str[last_dot_pos+1:]
            else:
                val_str = val_str.replace('.', '')
        
        try:
            return float(val_str)
        except Exception as e:
            print(f"Parse error for '{val}': {e}")
            return 0.0
    
    # Calculate statistics
    df_numeric = df_export.copy()
    df_numeric['Balance_numeric'] = df_numeric['Saldo'].apply(parse_amount)
    
    total_balance = df_numeric['Balance_numeric'].sum()
    avg_balance = df_numeric['Balance_numeric'].mean()
    highest_balance = df_numeric['Balance_numeric'].max()
    lowest_balance = df_numeric['Balance_numeric'].min()
    
    print(f"\nStatistics:")
    print(f"  Total: {total_balance:,.2f}")
    print(f"  Average: {avg_balance:,.2f}")
    print(f"  Highest: {highest_balance:,.2f}")
    print(f"  Lowest: {lowest_balance:,.2f}")
    
    # Add statistics to Excel
    last_row = len(df_export) + 2
    stats_start_row = last_row + 1
    
    worksheet.cell(row=stats_start_row, column=1, value='Total')
    worksheet.cell(row=stats_start_row, column=2, value=total_balance)
    worksheet.cell(row=stats_start_row, column=2).number_format = '#,##0.00'
    
    worksheet.cell(row=stats_start_row + 1, column=1, value='Rata-rata Pengendapan')
    worksheet.cell(row=stats_start_row + 1, column=2, value=avg_balance)
    worksheet.cell(row=stats_start_row + 1, column=2).number_format = '#,##0.00'
    
    worksheet.cell(row=stats_start_row + 2, column=1, value='Saldo Rata-rata')
    worksheet.cell(row=stats_start_row + 2, column=2, value=avg_balance)
    worksheet.cell(row=stats_start_row + 2, column=2).number_format = '#,##0.00'
    
    worksheet.cell(row=stats_start_row + 3, column=1, value='Saldo Tertinggi')
    worksheet.cell(row=stats_start_row + 3, column=2, value=highest_balance)
    worksheet.cell(row=stats_start_row + 3, column=2).number_format = '#,##0.00'
    
    worksheet.cell(row=stats_start_row + 4, column=1, value='Saldo Terendah')
    worksheet.cell(row=stats_start_row + 4, column=2, value=lowest_balance)
    worksheet.cell(row=stats_start_row + 4, column=2).number_format = '#,##0.00'
    
    # Mutation table
    mutation_start_row = stats_start_row + 6
    
    worksheet.merge_cells(start_row=mutation_start_row, start_column=1, end_row=mutation_start_row, end_column=3)
    header_cell = worksheet.cell(row=mutation_start_row, column=1)
    header_cell.value = 'Mutasi Debet'
    header_cell.font = Font(bold=True)
    header_cell.alignment = Alignment(horizontal='center')
    
    worksheet.merge_cells(start_row=mutation_start_row, start_column=4, end_row=mutation_start_row, end_column=6)
    header_cell2 = worksheet.cell(row=mutation_start_row, column=4)
    header_cell2.value = 'Mutasi kredit'
    header_cell2.font = Font(bold=True)
    header_cell2.alignment = Alignment(horizontal='center')
    
    mutation_rows = [
        ('Total Mutasi', '-', '-'),
        ('Adjusted', '-', '-'),
        ('Total Frekuensi', '-', '-'),
        ('Adjusted', '-', '-')
    ]
    
    thin_border = Border(
        left=Side(style='thin'),
        right=Side(style='thin'),
        top=Side(style='thin'),
        bottom=Side(style='thin')
    )
    
    for idx, (label, debit_val, credit_val) in enumerate(mutation_rows, start=1):
        row_num = mutation_start_row + idx
        worksheet.cell(row=row_num, column=1, value=label)
        worksheet.merge_cells(start_row=row_num, start_column=2, end_row=row_num, end_column=3)
        worksheet.cell(row=row_num, column=2, value=debit_val)
        worksheet.merge_cells(start_row=row_num, start_column=5, end_row=row_num, end_column=6)
        worksheet.cell(row=row_num, column=5, value=credit_val)
        
        for col in range(1, 7):
            worksheet.cell(row=row_num, column=col).border = thin_border
    
    for col in range(1, 7):
        worksheet.cell(row=mutation_start_row, column=col).border = thin_border
    
    # Column widths
    worksheet.column_dimensions['A'].width = 20
    worksheet.column_dimensions['B'].width = 20

print(f"\n✅ Excel file created: {output_path}")
print("\nPlease open the file to verify:")
print("  1. Daily balances appear correctly")
print("  2. Statistics are calculated")
print("  3. Mutation table is formatted")
print("  4. Day numbers (1-31) appear in Tanggal column")
