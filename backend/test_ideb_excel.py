"""Test IDEB SUYANTO Excel export"""
from processors.ideb_processor import process_ideb_file
import pandas as pd
from openpyxl import load_workbook
from openpyxl.styles import Font, Alignment, PatternFill, Border, Side
from openpyxl.utils import get_column_letter
import os

# Process the PDF
df = process_ideb_file('../test_data/IDEB SUYANTO.pdf', 'pdf')

# Export to Excel
output_path = 'outputs/test_ideb_suyanto.xlsx'
df.to_excel(output_path, index=False, sheet_name='Data Kredit')

print(f"✅ Excel exported to: {output_path}")

# Now enhance with formatting
wb = load_workbook(output_path)
ws = wb.active

# Style definitions
header_fill = PatternFill(start_color="366092", end_color="366092", fill_type="solid")
header_font = Font(bold=True, color="FFFFFF", size=11)
total_fill = PatternFill(start_color="DCE6F1", end_color="DCE6F1", fill_type="solid")
total_font = Font(bold=True, size=11)
border_side = Side(style='thin', color='000000')
border = Border(left=border_side, right=border_side, top=border_side, bottom=border_side)

# Format header row
for cell in ws[1]:
    cell.fill = header_fill
    cell.font = header_font
    cell.alignment = Alignment(horizontal='center', vertical='center')
    cell.border = border

# Format number columns (Plafon, O/S, Angsuran) with thousand separator
number_columns = ['B', 'D', 'I']  # Plafon, O/S, Angsuran
for col in number_columns:
    for row in range(2, ws.max_row + 1):
        cell = ws[f'{col}{row}']
        cell.number_format = '#,##0'  # Thousand separator, no decimals
        cell.alignment = Alignment(horizontal='right')
        cell.border = border

# Format other columns
for row in range(2, ws.max_row + 1):
    for col in ['A', 'C', 'E', 'F', 'G', 'H']:  # Text columns
        cell = ws[f'{col}{row}']
        cell.alignment = Alignment(horizontal='left' if col == 'A' else 'center')
        cell.border = border

# Add Total row
total_row = ws.max_row + 1
ws[f'A{total_row}'] = 'Total'
ws[f'A{total_row}'].font = total_font
ws[f'A{total_row}'].fill = total_fill
ws[f'A{total_row}'].border = border

# Add SUM formulas for numeric columns
ws[f'B{total_row}'] = f'=SUM(B2:B{total_row-1})'  # Plafon
ws[f'D{total_row}'] = f'=SUM(D2:D{total_row-1})'  # O/S
ws[f'I{total_row}'] = f'=SUM(I2:I{total_row-1})'  # Angsuran

# Format total row cells
for col in ['B', 'D', 'I']:
    cell = ws[f'{col}{total_row}']
    cell.font = total_font
    cell.fill = total_fill
    cell.number_format = '#,##0'
    cell.alignment = Alignment(horizontal='right')
    cell.border = border

# Empty cells in total row
for col in ['C', 'E', 'F', 'G', 'H']:
    cell = ws[f'{col}{total_row}']
    cell.fill = total_fill
    cell.border = border

# Adjust column widths
ws.column_dimensions['A'].width = 30  # Nama Bank
ws.column_dimensions['B'].width = 18  # Plafon
ws.column_dimensions['C'].width = 12  # Yield
ws.column_dimensions['D'].width = 18  # O/S
ws.column_dimensions['E'].width = 16  # Tanggal Pencairan
ws.column_dimensions['F'].width = 18  # Tanggal Jatuh Tempo
ws.column_dimensions['G'].width = 10  # Jk Waktu
ws.column_dimensions['H'].width = 8   # Kol
ws.column_dimensions['I'].width = 18  # Angsuran

# Save
wb.save(output_path)
print(f"✅ Excel formatted and saved!")
print(f"📂 Location: {os.path.abspath(output_path)}")
