"""Test full IDEB flow like in bot"""
from processors.ideb_processor import process_ideb_file
import pandas as pd
import os
from datetime import datetime

# Simulate bot export function (simplified)
def export_to_excel_test(df, filename):
    """Export DataFrame to Excel with professional formatting"""
    
    # Export to Excel
    df.to_excel(filename, index=False, sheet_name='Data Kredit')
    
    # Apply professional formatting
    try:
        from openpyxl import load_workbook
        from openpyxl.styles import Font, Alignment, PatternFill, Border, Side
        
        wb = load_workbook(filename)
        ws = wb.active
        
        # Style definitions
        header_fill = PatternFill(start_color="366092", end_color="366092", fill_type="solid")
        header_font = Font(bold=True, color="FFFFFF", size=11)
        total_fill = PatternFill(start_color="DCE6F1", end_color="DCE6F1", fill_type="solid")
        total_font = Font(bold=True, size=11)
        border_side = Side(style='thin', color='000000')
        border = Border(left=border_side, right=border_side, top=border_side, bottom=border_side)
        
        # Format header row
        for cell in ws[1]:
            cell.fill = header_fill
            cell.font = header_font
            cell.alignment = Alignment(horizontal='center', vertical='center')
            cell.border = border
        
        # Format number columns (Plafon, O/S, Angsuran) with thousand separator
        number_columns = ['B', 'D', 'I']  # Plafon, O/S, Angsuran
        for col in number_columns:
            for row in range(2, ws.max_row + 1):
                cell = ws[f'{col}{row}']
                cell.number_format = '#,##0'  # Thousand separator, no decimals
                cell.alignment = Alignment(horizontal='right')
                cell.border = border
        
        # Format other columns
        for row in range(2, ws.max_row + 1):
            for col in ['A', 'C', 'E', 'F', 'G', 'H']:  # Text columns
                cell = ws[f'{col}{row}']
                cell.alignment = Alignment(horizontal='left' if col == 'A' else 'center')
                cell.border = border
        
        # Add Total row
        total_row = ws.max_row + 1
        ws[f'A{total_row}'] = 'Total'
        ws[f'A{total_row}'].font = total_font
        ws[f'A{total_row}'].fill = total_fill
        ws[f'A{total_row}'].border = border
        
        # Add SUM formulas for numeric columns
        ws[f'B{total_row}'] = f'=SUM(B2:B{total_row-1})'  # Plafon
        ws[f'D{total_row}'] = f'=SUM(D2:D{total_row-1})'  # O/S
        ws[f'I{total_row}'] = f'=SUM(I2:I{total_row-1})'  # Angsuran
        
        # Format total row cells
        for col in ['B', 'D', 'I']:
            cell = ws[f'{col}{total_row}']
            cell.font = total_font
            cell.fill = total_fill
            cell.number_format = '#,##0'
            cell.alignment = Alignment(horizontal='right')
            cell.border = border
        
        # Empty cells in total row
        for col in ['C', 'E', 'F', 'G', 'H']:
            cell = ws[f'{col}{total_row}']
            cell.fill = total_fill
            cell.border = border
        
        # Adjust column widths
        ws.column_dimensions['A'].width = 30  # Nama Bank
        ws.column_dimensions['B'].width = 18  # Plafon
        ws.column_dimensions['C'].width = 12  # Yield
        ws.column_dimensions['D'].width = 18  # O/S
        ws.column_dimensions['E'].width = 16  # Tanggal Pencairan
        ws.column_dimensions['F'].width = 18  # Tanggal Jatuh Tempo
        ws.column_dimensions['G'].width = 10  # Jk Waktu
        ws.column_dimensions['H'].width = 8   # Kol
        ws.column_dimensions['I'].width = 18  # Angsuran
        
        # Save
        wb.save(filename)
        print("✅ Excel formatted with totals and professional styling")
    except Exception as e:
        print(f"⚠️ Could not apply formatting: {e}")

# Process IDEB SUYANTO
print("=" * 80)
print("TEST: Full IDEB Bot Flow")
print("=" * 80)

df = process_ideb_file('../test_data/IDEB SUYANTO.pdf', 'pdf')

if len(df) > 0:
    print(f"\n✅ Extracted {len(df)} credits")
    
    # Export like bot would
    output_file = 'outputs/ideb_slik_test_bot_flow.xlsx'
    export_to_excel_test(df, output_file)
    
    print(f"\n📂 File saved: {os.path.abspath(output_file)}")
    print("\n" + "=" * 80)
    print("Silakan buka file Excel dan verifikasi:")
    print("=" * 80)
    print("1. ✅ Data tanpa desimal ,00")
    print("2. ✅ Baris Total di bawah dengan formula SUM()")
    print("3. ✅ Total Plafon = 5,244,273,900")
    print("4. ✅ Total O/S = 4,451,629,582")
    print("5. ✅ Total Angsuran = 472,575,029")
    print("6. ✅ Header dengan background biru")
    print("7. ✅ Border pada semua cell")
else:
    print("❌ No data extracted")
